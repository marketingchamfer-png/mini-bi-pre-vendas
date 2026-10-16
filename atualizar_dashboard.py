# -*- coding: utf-8 -*-
"""
Mini BI de Pré-Vendas — Chamfer
================================

Lê as planilhas semanais de leads e vendas da pasta de rede, limpa, cruza e
gera um painel HTML único e autossuficiente (dados + bibliotecas embutidos).

Uso:
    python atualizar_dashboard.py                 # usa config.json ao lado
    python atualizar_dashboard.py --pasta-raiz X  # lê de outra pasta (testes)
    python atualizar_dashboard.py --saida Y       # grava em outra pasta

As planilhas de origem são abertas SOMENTE para leitura (a partir de uma
cópia temporária). Se algo falhar, o dashboard.html anterior é mantido.
"""

from __future__ import annotations

import argparse
import base64
import datetime as dt
import fnmatch
import json
import logging
import os
import re
import shutil
import sys
import tempfile
import time
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

PASTA_APP = Path(__file__).resolve().parent
PASTA_RECURSOS = PASTA_APP / "recursos"

# Siglas e nomes das UFs (geografia fixa; as regiões ficam no config.json)
UFS = {
    "AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas", "BA": "Bahia",
    "CE": "Ceará", "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás",
    "MA": "Maranhão", "MT": "Mato Grosso", "MS": "Mato Grosso do Sul", "MG": "Minas Gerais",
    "PA": "Pará", "PB": "Paraíba", "PR": "Paraná", "PE": "Pernambuco", "PI": "Piauí",
    "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte", "RS": "Rio Grande do Sul",
    "RO": "Rondônia", "RR": "Roraima", "SC": "Santa Catarina", "SP": "São Paulo",
    "SE": "Sergipe", "TO": "Tocantins",
}

log = logging.getLogger("mini_bi")


class ErroAtualizacao(Exception):
    """Erro esperado, com mensagem em português para o usuário."""


# ---------------------------------------------------------------------------
# Texto e números
# ---------------------------------------------------------------------------

def sem_acentos(texto: str) -> str:
    """Remove acentos mantendo as letras base (ç → c, ã → a)."""
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def limpar(valor) -> str:
    """Converte para texto, apara e colapsa espaços internos. None → ''."""
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def chave(texto) -> str:
    """Chave de comparação: maiúsculas, sem acentos, só letras e números."""
    return re.sub(r"[^A-Z0-9]", "", sem_acentos(limpar(texto)).upper())


def normalizar_empresa(nome) -> str:
    """Maiúsculas, sem acentos, sem pontuação, espaços simples (mantém números)."""
    s = sem_acentos(limpar(nome)).upper()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def numero(valor):
    """Converte número em formato brasileiro ('R$ 1.234,56', '5.000') para float."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    s = limpar(valor).replace("R$", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def numeros_no_texto(texto) -> list[float]:
    """Extrai todos os números de um texto ('5.000 E 1.260' → [5000, 1260])."""
    if isinstance(texto, (int, float)) and not isinstance(texto, bool):
        return [float(texto)]
    achados = re.findall(r"\d[\d.,]*\d|\d", limpar(texto))
    return [n for n in (numero(a) for a in achados) if n is not None]


def para_data(valor):
    """Aceita datetime, date, serial do Excel ou texto dd/mm/aaaa [hh:mm[:ss]]."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, dt.datetime):
        return valor
    if isinstance(valor, dt.date):
        return dt.datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(valor))
    s = limpar(valor)
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def iso(d):
    return d.strftime("%Y-%m-%dT%H:%M:%S") if d else None


# ---------------------------------------------------------------------------
# Configuração e seleção de arquivos
# ---------------------------------------------------------------------------

def carregar_config(caminho: Path) -> dict:
    if not caminho.exists():
        raise ErroAtualizacao(f"Arquivo de configuração não encontrado: {caminho}")
    try:
        return json.loads(caminho.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ErroAtualizacao(
            f"O config.json tem um erro de formatação (linha {e.lineno}, coluna {e.colno}): {e.msg}")


def escolher_arquivos(cfg: dict, raiz: Path) -> dict:
    """Para cada padrão de nome, escolhe o arquivo mais recente (empate: nome maior)."""
    if not raiz.exists():
        raise ErroAtualizacao(
            f"A pasta de origem não foi encontrada: {raiz}\n"
            "Verifique se a unidade de rede Y: está conectada (abra o Explorador de Arquivos "
            "e clique na unidade Y:) e tente novamente. O dashboard anterior foi mantido.")
    extensoes = {e.lower() for e in cfg["extensoes_aceitas"]}
    ignoradas = {chave(p) for p in cfg.get("subpastas_ignoradas", [])}
    iterador = raiz.rglob("*") if cfg.get("varrer_subpastas", True) else raiz.glob("*")
    candidatos = []
    for p in iterador:
        if not p.is_file() or p.name.startswith("~$") or p.suffix.lower() not in extensoes:
            continue
        if any(chave(parte) in ignoradas for parte in p.relative_to(raiz).parts[:-1]):
            continue
        candidatos.append(p)

    escolhidos = {}
    for tipo, padrao in cfg["padroes_de_nome"].items():
        padrao_norm = sem_acentos(padrao).upper()
        casam = [p for p in candidatos
                 if fnmatch.fnmatchcase(sem_acentos(p.stem).upper(), padrao_norm)]
        if not casam:
            raise ErroAtualizacao(
                f"Nenhuma planilha encontrada para '{tipo}' (padrão {padrao}) em {raiz}.")
        casam.sort(key=lambda p: (p.stat().st_mtime, p.name.upper()), reverse=True)
        escolhidos[tipo] = casam[0]
        if len(casam) > 1:
            log.info("  %s: %d versões encontradas; usando a mais recente.", tipo, len(casam))
    return escolhidos


def copiar_para_temporario(origem: Path, pasta_tmp: Path, tentativas: int = 4) -> Path:
    """Copia a planilha antes de ler (evita bloqueio quando alguém está com ela aberta)."""
    destino = pasta_tmp / f"{origem.stem}_{int(time.time() * 1000)}{origem.suffix}"
    ultimo_erro = None
    for n in range(1, tentativas + 1):
        try:
            shutil.copy2(origem, destino)
            return destino
        except OSError as e:
            ultimo_erro = e
            log.warning("  Arquivo em uso ou inacessível (%s), tentativa %d/%d...",
                        origem.name, n, tentativas)
            time.sleep(3)
    raise ErroAtualizacao(f"Não foi possível ler {origem.name}: {ultimo_erro}")


# ---------------------------------------------------------------------------
# Leitura bruta (openpyxl, somente leitura)
# ---------------------------------------------------------------------------

def ler_aba(caminho: Path, nome_aba: str, limite_vazias: int) -> tuple[list, list, list[str]]:
    """Devolve (cabeçalho, linhas, nomes_das_abas). Para após N linhas vazias seguidas."""
    import openpyxl  # importado aqui para a mensagem de erro de dependência ficar clara

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        abas = list(wb.sheetnames)
        alvo = next((a for a in abas if chave(a) == chave(nome_aba)), None)
        if alvo is None:
            raise ErroAtualizacao(
                f"A aba '{nome_aba}' não existe em {caminho.name}. Abas encontradas: {', '.join(abas)}")
        linhas, vazias, cabecalho = [], 0, None
        for linha in wb[alvo].iter_rows(values_only=True):
            if cabecalho is None:
                cabecalho = list(linha)
                continue
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in linha):
                vazias += 1
                if vazias >= limite_vazias:
                    break
                continue
            vazias = 0
            linhas.append(list(linha))
        return cabecalho or [], linhas, abas
    finally:
        wb.close()


def serializar_celula(v):
    """Datas viram texto ISO para o JSON bruto (mesmo formato que o navegador gera)."""
    if isinstance(v, (dt.datetime, dt.date)):
        return iso(para_data(v))
    return v


# ---------------------------------------------------------------------------
# Transformação (espelhada em recursos/painel.js → ETL.transformar)
# ---------------------------------------------------------------------------

class Mapeador:
    """Classifica status e itens conforme os mapas do config.json."""

    def __init__(self, cfg: dict):
        ms = cfg["mapa_de_status"]
        self.status_exato = {}
        for grupo, lista in ms["grupos"].items():
            for s in lista:
                self.status_exato[chave(s)] = (grupo, s)
        self.motivo_exato = {chave(s): motivo
                             for motivo, lista in ms["motivos_de_perda"].items() for s in lista}
        self.regras_contem = [(chave(r["contem"]), r) for r in ms.get("regras_contem", [])]

        mi = cfg["mapa_de_itens"]
        self.de_para = {chave(k): v for k, v in mi["de_para"].items()}
        self.revisar = {chave(k): v for k, v in mi.get("revisar", {}).items()}
        self.sep_itens = re.compile(mi["separadores_de_itens"], re.IGNORECASE)
        self.familias = [(re.compile(f["regex"]), f["familia"]) for f in mi["familias"]]
        self.familia_padrao = mi["familia_padrao"]

    def status(self, original: str):
        """Devolve (status_padrao, grupo, motivo). Grupo None = não mapeado."""
        k = chave(original)
        if not k:
            return "", None, None
        if k in self.status_exato:
            grupo, padrao = self.status_exato[k]
            return padrao, grupo, self.motivo_exato.get(k)
        # Textos longos: se contiver um status conhecido, herda o grupo do mais longo
        contidos = [s for s in self.status_exato if len(s) >= 10 and s in k]
        if contidos:
            s = max(contidos, key=len)
            grupo, padrao = self.status_exato[s]
            return limpar(original), grupo, self.motivo_exato.get(s)
        for termo, regra in self.regras_contem:
            if termo and termo in k:
                return limpar(original), regra["grupo"], regra.get("motivo")
        return limpar(original), None, None

    def separar_itens(self, texto: str) -> list[str]:
        return [p for p in (limpar(x) for x in self.sep_itens.split(limpar(texto))) if p]

    def padronizar_item(self, item: str) -> str:
        k = chave(item)
        if k in self.de_para:
            return self.de_para[k]
        s = limpar(item)
        if re.fullmatch(r"\d+\.\d+[A-Za-z]*", s):          # 1.8 → 1,8
            s = s.replace(".", ",")
        if re.fullmatch(r"[Pp]?\d+(,\d+)?\s*[A-Za-z.]*", s):  # códigos: maiúsculas, sem espaços/pontos
            s = re.sub(r"[\s.]", "", s).upper()
        return s

    def familia(self, item_padrao: str) -> str:
        alvo = sem_acentos(item_padrao).upper()
        for rx, fam in self.familias:
            if rx.search(alvo):
                return fam
        return self.familia_padrao


def indice_colunas(cabecalho: list) -> list[tuple[int, str]]:
    return [(i, limpar(h)) for i, h in enumerate(cabecalho) if limpar(h)]


def achar_coluna(colunas, nome, avisos, arquivo, obrigatoria=False):
    alvo = chave(nome)
    for i, h in colunas:
        if chave(h) == alvo:
            return i
    # tolerância: cabeçalho que começa com o nome (ex.: 'Nome da empresa')
    for i, h in colunas:
        if chave(h).startswith(alvo[:12]):
            return i
    msg = f"Coluna '{nome}' não encontrada em {arquivo}."
    if obrigatoria:
        raise ErroAtualizacao(msg)
    avisos.append(msg)
    return None


def colunas_de_pedidos(colunas, cfg, ano_base, mes_base, avisos):
    """Colunas 'Pedidos <mês> [ano]' → [(indice, 'AAAA-MM', cabeçalho)]."""
    meses = {chave(k): v for k, v in cfg["meses"].items()}
    saida = []
    for i, h in colunas:
        k = chave(h)
        if not k.startswith("PEDIDOS"):
            continue
        resto = k[len("PEDIDOS"):]
        ano = None
        m_ano = re.search(r"(20\d\d)$", resto)
        if m_ano:
            ano = int(m_ano.group(1))
            resto = resto[: m_ano.start()]
        mes = meses.get(resto)
        if mes is None:
            avisos.append(f"Coluna de pedidos com mês não reconhecido: '{h}' (ignorada).")
            continue
        if ano is None:  # sem ano no cabeçalho: mês posterior ao último lead pertence ao ano anterior
            ano = ano_base if mes <= mes_base else ano_base - 1
        saida.append((i, f"{ano:04d}-{mes:02d}", h))
    return saida


def nome_vendedor_base(nome: str, sufixos: list[str]) -> str:
    s = normalizar_empresa(nome)
    for suf in sufixos:
        s = re.sub(rf"\b{re.escape(suf)}\b", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def transformar(bruto: dict, cfg: dict) -> dict:
    """Recebe as linhas brutas das duas planilhas e devolve o modelo de dados do painel."""
    mp = Mapeador(cfg)
    avisos: list[str] = []
    qualidade: list[dict] = []

    # ---------- Planilha de status / vendas ----------
    st = bruto["status_vendas"]
    cols = indice_colunas(st["cabecalho"])
    arq_s = st["arquivo"]
    c_data = achar_coluna(cols, "Data", avisos, arq_s)
    c_vend = achar_coluna(cols, "Vendedor", avisos, arq_s)
    c_emp = achar_coluna(cols, "Nome da empresa (Razão Social)", avisos, arq_s, obrigatoria=True)
    c_score = achar_coluna(cols, "Score", avisos, arq_s)
    c_status = achar_coluna(cols, "Status", avisos, arq_s)

    def val(linha, i):
        return linha[i] if i is not None and i < len(linha) else None

    usadas_s = {c for c in (c_data, c_vend, c_emp, c_score, c_status) if c is not None}
    datas_s = [para_data(val(l, c_data)) for l in st["linhas"]]
    datas_validas = [d for d in datas_s if d]
    ref = max(datas_validas) if datas_validas else dt.datetime.now()
    pedidos_cols = colunas_de_pedidos(cols, cfg, ref.year, ref.month, avisos)
    usadas_s |= {i for i, _, _ in pedidos_cols}

    linhas_s = []
    for n, l in enumerate(st["linhas"]):
        if all(val(l, i) is None or limpar(val(l, i)) == "" for i in usadas_s):
            continue  # linha só com legendas (colunas à direita) ou vazia
        empresa = limpar(val(l, c_emp))
        if not empresa:
            continue
        score_txt = limpar(val(l, c_score))
        m = re.search(r"(\d+)", score_txt)
        score = int(m.group(1)) if m and int(m.group(1)) in (1, 2, 3) else None
        pedidos = []
        for i, ano_mes, h in pedidos_cols:
            bruto_v = val(l, i)
            if bruto_v is None or limpar(bruto_v) == "":
                continue
            v = numero(bruto_v)
            if v is None:
                qualidade.append({"tipo_problema": "Valor de pedido não numérico", "lead_id": None,
                                  "empresa": empresa, "vendedor": limpar(val(l, c_vend)),
                                  "detalhe": f"{h}: '{limpar(bruto_v)}'"})
                continue
            if v != 0:
                pedidos.append((ano_mes, v))
        linhas_s.append({
            "linha": n + 2, "data": datas_s[n], "vendedor": limpar(val(l, c_vend)),
            "empresa": empresa, "norm": normalizar_empresa(empresa),
            "score": score, "score_txt": score_txt,
            "status": limpar(val(l, c_status)), "pedidos": pedidos,
        })

    # ---------- Planilha de potencial ----------
    pt = bruto["potencial_leads"]
    colp = indice_colunas(pt["cabecalho"])
    arq_p = pt["arquivo"]
    p_data = achar_coluna(colp, "Data", avisos, arq_p)
    p_vend = achar_coluna(colp, "Vendedor", avisos, arq_p)
    p_emp = achar_coluna(colp, "Nome da empresa (Razão Social)", avisos, arq_p, obrigatoria=True)
    p_uf = achar_coluna(colp, "Estado", avisos, arq_p)
    p_val = achar_coluna(colp, "Valor total", avisos, arq_p)
    # Pares item NN + Quantidade, detectados pelo nome (não pela letra da coluna)
    pares = []
    for pos, (i, h) in enumerate(colp):
        if re.fullmatch(r"ITEM\d*", chave(h)):
            prox = colp[pos + 1] if pos + 1 < len(colp) else None
            q = prox[0] if prox and chave(prox[1]).startswith("QUANTIDADE") else None
            pares.append((i, q))
    if not pares:
        avisos.append(f"Nenhuma coluna 'item NN' encontrada em {arq_p}.")
    usadas_p = {c for c in (p_data, p_vend, p_emp, p_uf, p_val) if c is not None}
    usadas_p |= {i for par in pares for i in par if i is not None}

    linhas_p = []
    for n, l in enumerate(pt["linhas"]):
        if all(val(l, i) is None or limpar(val(l, i)) == "" for i in usadas_p):
            continue
        empresa = limpar(val(l, p_emp))
        if not empresa:
            continue
        itens = []
        for ci, cq in pares:
            item_txt = limpar(val(l, ci))
            qtd_bruta = val(l, cq) if cq is not None else None
            if not item_txt:
                if qtd_bruta not in (None, "") and limpar(qtd_bruta):
                    itens.append({"erro": "Quantidade sem item", "detalhe": limpar(qtd_bruta)})
                continue
            codigos = mp.separar_itens(item_txt)
            qtds = numeros_no_texto(qtd_bruta) if qtd_bruta not in (None, "") else []
            alerta = None
            if qtd_bruta not in (None, "") and limpar(qtd_bruta) and not qtds:
                alerta = f"Quantidade não numérica: '{limpar(qtd_bruta)}' ({item_txt})"
            if len(codigos) == len(qtds):
                pares_q = list(zip(codigos, qtds))
            elif len(qtds) == 1:
                pares_q = [(c, qtds[0]) for c in codigos]
                if len(codigos) > 1:
                    alerta = f"Uma quantidade para vários itens: '{item_txt}' = {limpar(qtd_bruta)}"
            elif len(qtds) > 1 and len(codigos) == 1:
                pares_q = [(codigos[0], sum(qtds))]
                alerta = f"Quantidade composta somada: '{limpar(qtd_bruta)}' ({item_txt})"
            else:
                pares_q = [(c, None) for c in codigos]
                if qtds:
                    alerta = f"Quantidades não batem com os itens: '{item_txt}' = {limpar(qtd_bruta)}"
            for cod, q in pares_q:
                padrao = mp.padronizar_item(cod)
                itens.append({"item_original": cod, "item_padrao": padrao,
                              "familia": mp.familia(padrao), "quantidade": q})
            if alerta:
                itens.append({"erro": "Quantidade a revisar", "detalhe": alerta})
        uf_txt = limpar(val(l, p_uf))
        valor_bruto = val(l, p_val)
        linhas_p.append({
            "linha": n + 2, "data": para_data(val(l, p_data)), "vendedor": limpar(val(l, p_vend)),
            "empresa": empresa, "norm": normalizar_empresa(empresa), "estado": uf_txt,
            "potencial": numero(valor_bruto),
            "potencial_invalido": valor_bruto not in (None, "") and numero(valor_bruto) is None,
            "itens": itens,
        })

    # ---------- Junção (full outer) ----------
    sufixos = cfg["vendedores"]["sufixos_ignorados"]

    def dia(d):
        return d.date().isoformat() if d else ""

    por_chave_p = defaultdict(list)
    for j, p in enumerate(linhas_p):
        por_chave_p[(p["norm"], dia(p["data"]))].append(j)
    par_de_s: dict[int, int] = {}
    usados_p: set[int] = set()
    por_chave_s = defaultdict(list)
    for i, s in enumerate(linhas_s):
        por_chave_s[(s["norm"], dia(s["data"]))].append(i)
    for k, lista_s in por_chave_s.items():
        livres = [j for j in por_chave_p.get(k, []) if j not in usados_p]
        # primeiro casa quem tem o mesmo vendedor, depois pela ordem
        for i in list(lista_s):
            vb = nome_vendedor_base(linhas_s[i]["vendedor"], sufixos)
            j = next((j for j in livres if nome_vendedor_base(linhas_p[j]["vendedor"], sufixos) == vb), None)
            if j is not None:
                par_de_s[i] = j
                livres.remove(j)
                usados_p.add(j)
        for i in lista_s:
            if i not in par_de_s and livres:
                j = livres.pop(0)
                par_de_s[i] = j
                usados_p.add(j)
    casados_chave = len(par_de_s)

    # Fallback: empresa única nos dois arquivos
    cont_s = Counter(s["norm"] for s in linhas_s)
    cont_p = Counter(p["norm"] for p in linhas_p)
    idx_p_por_norm = {p["norm"]: j for j, p in enumerate(linhas_p)}
    tipo_corresp = {i: "empresa+data" for i in par_de_s}
    for i, s in enumerate(linhas_s):
        if i in par_de_s:
            continue
        if cont_s[s["norm"]] == 1 and cont_p.get(s["norm"]) == 1:
            j = idx_p_por_norm[s["norm"]]
            if j not in usados_p:
                par_de_s[i] = j
                usados_p.add(j)
                tipo_corresp[i] = "empresa"

    registros = []
    for i, s in enumerate(linhas_s):
        j = par_de_s.get(i)
        registros.append((s, linhas_p[j] if j is not None else None, tipo_corresp.get(i, "nenhuma")))
    for j, p in enumerate(linhas_p):
        if j not in usados_p:
            registros.append((None, p, "nenhuma"))

    def ordem(r):
        s, p, _ = r
        d = (s or p)["data"]
        return (d or dt.datetime.max, (s or p)["norm"])
    registros.sort(key=ordem)

    # ---------- Modelo final ----------
    regioes = {uf: reg for reg, lista in cfg["regioes"].items() for uf in lista}
    nomes_uf = {chave(nome): sig for sig, nome in UFS.items()}
    limite = float(cfg["limite_outlier_reais"])
    leads, pedidos, itens = [], [], []
    status_nao_mapeados, itens_revisar, ufs_nao_mapeadas = Counter(), {}, Counter()

    def flag(lead, tipo, detalhe=""):
        lead["flags_qualidade"].append(tipo)
        qualidade.append({"tipo_problema": tipo, "lead_id": lead["lead_id"], "empresa": lead["empresa"],
                          "vendedor": lead["vendedor"], "detalhe": detalhe})

    for n, (s, p, corresp) in enumerate(registros, start=1):
        base = s or p
        lead_id = f"L{n:05d}"
        data = base["data"] or (p["data"] if p else None)
        status_orig = s["status"] if s else ""
        status_padrao, grupo, motivo = mp.status(status_orig) if s else ("", None, None)
        if s and grupo is None:
            grupo = "Outros (revisar)"
            status_nao_mapeados[status_orig or "(vazio)"] += 1
        elif not s:
            grupo = "Sem pedido registrado"
            status_padrao = "(sem linha na planilha de status)"
        # UF
        uf_sig, uf_nome = None, None
        if p and p["estado"]:
            partes = [x.strip() for x in p["estado"].rsplit("-", 1)]
            cand = chave(partes[-1]) if partes else ""
            if cand in UFS:
                uf_sig = cand
            elif chave(p["estado"]) in nomes_uf:
                uf_sig = nomes_uf[chave(p["estado"])]
            elif chave(partes[0]) in nomes_uf:
                uf_sig = nomes_uf[chave(partes[0])]
            if uf_sig:
                uf_nome = UFS[uf_sig]
            else:
                ufs_nao_mapeadas[p["estado"]] += 1
        faturado = round(sum(v for _, v in (s["pedidos"] if s else [])), 2)
        potencial = p["potencial"] if p else None
        lead = {
            "lead_id": lead_id,
            "data_entrada": iso(data),
            "mes_entrada": data.strftime("%Y-%m") if data else None,
            "vendedor": (s["vendedor"] if s else "") or (p["vendedor"] if p else ""),
            "empresa": base["empresa"],
            "empresa_normalizada": base["norm"],
            "uf_sigla": uf_sig,
            "uf_nome": uf_nome,
            "regiao": regioes.get(uf_sig) if uf_sig else None,
            "score": s["score"] if s else None,
            "status_original": status_orig,
            "status_padrao": status_padrao,
            "status_grupo": grupo,
            "status_motivo_perda": motivo,
            "potencial_r$": potencial,
            "faturado_r$": faturado,
            "flag_outlier": bool(potencial is not None and potencial > limite),
            "flag_recorrente": False,
            "flag_sem_correspondencia": corresp == "nenhuma",
            "correspondencia": corresp,
            "origem": "ambos" if s and p else ("so_status" if s else "so_potencial"),
            "familias": [],
            "flags_qualidade": [],
        }
        leads.append(lead)

        for ano_mes, v in (s["pedidos"] if s else []):
            pedidos.append({"lead_id": lead_id, "ano_mes": ano_mes, "valor_r$": round(v, 2)})
        familias = []
        for it in (p["itens"] if p else []):
            if "erro" in it:
                flag(lead, it["erro"], it["detalhe"])
                continue
            itens.append({"lead_id": lead_id, **it})
            if it["familia"] not in familias:
                familias.append(it["familia"])
            k = chave(it["item_original"])
            if k in mp.revisar or it["familia"] == mp.familia_padrao:
                reg = itens_revisar.setdefault(it["item_original"], {
                    "item": it["item_original"], "item_padrao": it["item_padrao"],
                    "familia": it["familia"], "contagem": 0,
                    "sugestao": mp.revisar.get(k, "Sem regra de família: incluir no config.json")})
                reg["contagem"] += 1
        lead["familias"] = familias

        # Inconsistências
        if not data:
            flag(lead, "Data ausente ou inválida")
        if grupo == "Outros (revisar)":
            flag(lead, "Status não mapeado", status_orig or "(vazio)")
        if grupo == "Fechado" and faturado == 0:
            flag(lead, "Pedido fechado sem valor de pedido lançado")
        if faturado > 0 and grupo != "Fechado":
            flag(lead, "Valor de pedido lançado com status diferente de fechado",
                 f"Status: {status_orig or '(vazio)'}")
        if p is not None and potencial is None:
            flag(lead, "Sem valor de potencial",
                 "Valor não numérico" if p.get("potencial_invalido") else "Célula vazia")
        if corresp == "nenhuma":
            flag(lead, "Empresa só na planilha de status" if s else "Empresa só na planilha de potencial")
        elif corresp == "empresa":
            flag(lead, "Correspondência apenas pelo nome da empresa",
                 f"Datas: {dia(s['data'])} × {dia(p['data'])}")
        if s and p and nome_vendedor_base(s["vendedor"], sufixos) != nome_vendedor_base(p["vendedor"], sufixos):
            flag(lead, "Vendedor diferente entre as planilhas", f"{s['vendedor']} × {p['vendedor']}")
        if p and p["estado"] and not uf_sig:
            flag(lead, "UF não reconhecida", p["estado"])

    # Recorrentes e duplicadas
    cont_emp = Counter(l["empresa_normalizada"] for l in leads)
    vistos = Counter()
    for l in leads:
        l["flag_recorrente"] = cont_emp[l["empresa_normalizada"]] > 1
        k = (l["empresa_normalizada"], (l["data_entrada"] or "")[:10], nome_vendedor_base(l["vendedor"], sufixos))
        vistos[k] += 1
        if vistos[k] == 2:
            flag(l, "Linha duplicada (mesma empresa, data e vendedor)")
        elif vistos[k] > 2:
            flag(l, "Linha duplicada (mesma empresa, data e vendedor)")

    # Vendedores inativos ou sem aba própria
    abas = [a for a in bruto["status_vendas"].get("abas", []) if chave(a) != chave(cfg["abas"]["status_vendas"])]
    palavra_inativo = chave(cfg["vendedores"]["palavra_inativo"])
    base_abas = {nome_vendedor_base(a, sufixos): a for a in abas}
    inativos = {nome_vendedor_base(a, sufixos) for a in abas if palavra_inativo in chave(a)}
    leads_por_vendedor = Counter(l["vendedor"] for l in leads if l["vendedor"])
    vendedores_info = []
    for v, qtd in sorted(leads_por_vendedor.items()):
        vb = nome_vendedor_base(v, sufixos)
        aba = base_abas.get(vb)
        situacao = "ativo"
        if vb in inativos:
            situacao = "inativo"
            qualidade.append({"tipo_problema": "Vendedor marcado como inativo", "lead_id": None,
                              "empresa": "", "vendedor": v, "detalhe": f"Aba '{aba}' · {qtd} lead(s)"})
        elif aba is None:
            situacao = "sem aba"
            qualidade.append({"tipo_problema": "Vendedor sem aba própria", "lead_id": None,
                              "empresa": "", "vendedor": v, "detalhe": f"{qtd} lead(s) na aba Geral"})
        vendedores_info.append({"vendedor": v, "aba": aba, "situacao": situacao, "leads": qtd})

    datas = [l["data_entrada"] for l in leads if l["data_entrada"]]
    mes_parcial = max(datas)[:7] if datas else None
    resumo = {
        "linhas_status": len(st["linhas"]), "linhas_potencial": len(pt["linhas"]),
        "validas_status": len(linhas_s), "validas_potencial": len(linhas_p),
        "casados_empresa_data": casados_chave,
        "casados_so_empresa": sum(1 for t in tipo_corresp.values() if t == "empresa"),
        "so_status": sum(1 for l in leads if l["origem"] == "so_status"),
        "so_potencial": sum(1 for l in leads if l["origem"] == "so_potencial"),
        "empresas_em_ambos": len({s["norm"] for s in linhas_s} & {p["norm"] for p in linhas_p}),
    }
    return {
        "leads": leads,
        "pedidos_mensais": pedidos,
        "itens": itens,
        "qualidade": qualidade,
        "nao_mapeados": {
            "status": [{"valor": k, "contagem": v} for k, v in status_nao_mapeados.most_common()],
            "itens": sorted(itens_revisar.values(), key=lambda r: -r["contagem"]),
            "ufs": [{"valor": k, "contagem": v} for k, v in ufs_nao_mapeadas.most_common()],
        },
        "vendedores": vendedores_info,
        "meta": {
            "total_linhas_lidas": len(st["linhas"]) + len(pt["linhas"]),
            "total_linhas_validas": len(linhas_s) + len(linhas_p),
            "mes_parcial": mes_parcial,
            "data_min": min(datas) if datas else None,
            "data_max": max(datas) if datas else None,
            "meses_pedidos": sorted({am for _, am, _ in pedidos_cols}),
            "colunas_pedidos": [h for _, _, h in pedidos_cols],
            "resumo": resumo,
            "avisos": avisos,
        },
    }


# ---------------------------------------------------------------------------
# Validação contra os valores de referência
# ---------------------------------------------------------------------------

def validar(dados: dict, cfg: dict) -> list[dict]:
    import pandas as pd

    ref = cfg.get("valores_de_referencia") or {}
    if not ref:
        return []
    leads = pd.DataFrame(dados["leads"])
    ped = pd.DataFrame(dados["pedidos_mensais"]) if dados["pedidos_mensais"] else pd.DataFrame(columns=["ano_mes", "valor_r$"])
    so_s = leads[leads["origem"] != "so_potencial"]
    so_p = leads[leads["origem"] != "so_status"]
    r = dados["meta"]["resumo"]
    obtidos = {
        "leads_aba_geral": len(so_s),
        "pedidos_fechados": int((so_s["status_grupo"] == "Fechado").sum()),
        "faturado_total": round(float(leads["faturado_r$"].sum()), 2),
        "leads_potencial": len(so_p),
        "potencial_total": round(float(so_p["potencial_r$"].fillna(0).sum()), 2),
        "maior_lead": float(so_p["potencial_r$"].max() or 0),
        "leads_sp": int((so_p["uf_sigla"] == "SP").sum()),
        "empresas_em_ambos": r["empresas_em_ambos"],
        "chave_empresa_data_casando": r["casados_empresa_data"],
    }
    por_mes = ped.groupby("ano_mes")["valor_r$"].sum().round(2).to_dict()
    por_score = so_s["score"].fillna(0).astype(int).astype(str).value_counts().to_dict()
    saida = []
    for k, esperado in ref.items():
        if k.startswith("_"):
            continue
        if isinstance(esperado, dict):
            fonte = por_mes if k == "faturado_por_mes" else por_score
            for sub, e in esperado.items():
                o = fonte.get(sub, 0)
                saida.append({"indicador": f"{k}[{sub}]", "esperado": e, "obtido": o,
                              "ok": abs(float(o) - float(e)) < 0.01})
        else:
            o = obtidos.get(k)
            saida.append({"indicador": k, "esperado": esperado, "obtido": o,
                          "ok": o is not None and abs(float(o) - float(esperado)) < 0.01})
    return saida


# ---------------------------------------------------------------------------
# Geração do HTML
# ---------------------------------------------------------------------------

def json_para_script(obj) -> str:
    """JSON seguro para colocar dentro de <script> (sem '</' nem separadores de linha JS)."""
    s = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
    return s.replace("</", "<\\/").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")


def css_fontes() -> str:
    regras = []
    for arq in sorted((PASTA_RECURSOS / "fontes").glob("*.woff2")):
        m = re.match(r"(inter|montserrat)-latin-(\d+)-normal", arq.stem)
        if not m:
            continue
        familia = "Inter" if m.group(1) == "inter" else "Montserrat"
        b64 = base64.b64encode(arq.read_bytes()).decode("ascii")
        regras.append(
            f"@font-face{{font-family:'{familia}';font-style:normal;font-weight:{m.group(2)};"
            f"font-display:swap;src:url(data:font/woff2;base64,{b64}) format('woff2');}}")
    return "\n".join(regras)


def config_para_painel(cfg: dict) -> dict:
    chaves = ["versao_painel", "titulo_painel", "limite_outlier_reais", "cores_da_marca", "meses",
              "mapa_de_status", "mapa_de_itens", "vendedores", "regioes", "padroes_de_nome", "abas"]
    saida = {k: cfg[k] for k in chaves if k in cfg}
    saida["ufs"] = UFS
    return saida


def montar_html(dados: dict, cfg: dict) -> str:
    modelo = (PASTA_RECURSOS / "painel.html").read_text(encoding="utf-8")
    app = (PASTA_RECURSOS / "painel.js").read_text(encoding="utf-8")
    chart = (PASTA_RECURSOS / "chart.umd.min.js").read_text(encoding="utf-8")
    sheetjs = base64.b64encode((PASTA_RECURSOS / "xlsx.full.min.js").read_bytes()).decode("ascii")
    trocas = {
        "/*__FONTES__*/": css_fontes(),
        "/*__CHARTJS__*/": chart.replace("</script", "<\\/script"),
        "__SHEETJS_B64__": sheetjs,
        "/*__DADOS__*/null": json_para_script(dados),
        "/*__CONFIG__*/null": json_para_script(config_para_painel(cfg)),
        "/*__APP__*/": app.replace("</script", "<\\/script"),
        "__LOGO_PNG_B64__": base64.b64encode((PASTA_RECURSOS / "logo_chamfer.png").read_bytes()).decode("ascii"),
    }
    for marcador, conteudo in trocas.items():
        if marcador not in modelo:
            raise ErroAtualizacao(f"Modelo do painel corrompido: marcador {marcador} ausente.")
        modelo = modelo.replace(marcador, conteudo)
    return modelo


def gravar_atomico(caminho: Path, conteudo: str):
    """Grava em arquivo temporário e troca no fim — nunca deixa um painel pela metade."""
    tmp = caminho.with_suffix(caminho.suffix + ".tmp")
    tmp.write_text(conteudo, encoding="utf-8")
    os.replace(tmp, caminho)


# ---------------------------------------------------------------------------
# Execução
# ---------------------------------------------------------------------------

def configurar_log(pasta_saida: Path) -> Path:
    log.setLevel(logging.INFO)
    log.handlers.clear()
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s %(message)s", "%d/%m/%Y %H:%M:%S")
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(logging.Formatter("%(message)s"))
    log.addHandler(console)
    destino = None
    for pasta in (pasta_saida, PASTA_APP, Path(tempfile.gettempdir())):
        try:
            pasta.mkdir(parents=True, exist_ok=True)
            destino = pasta / "atualizacao.log"
            if destino.exists() and destino.stat().st_size > 2_000_000:  # mantém o log enxuto
                destino.replace(destino.with_suffix(".log.old"))
            arq = logging.FileHandler(destino, encoding="utf-8")
            arq.setFormatter(fmt)
            log.addHandler(arq)
            break
        except OSError:
            continue
    return destino


def executar(args) -> int:
    cfg = carregar_config(Path(args.config))
    raiz = Path(args.pasta_raiz or cfg["pasta_raiz"])
    saida_cfg = args.saida or cfg.get("pasta_de_saida") or ""
    pasta_saida = Path(saida_cfg) if saida_cfg else raiz / "BI"
    destino_log = configurar_log(pasta_saida if raiz.exists() or saida_cfg else PASTA_APP)

    log.info("=" * 70)
    log.info("Atualização do Mini BI de Pré-Vendas — início")
    log.info("Pasta de origem: %s", raiz)
    inicio = time.time()

    arquivos = escolher_arquivos(cfg, raiz)
    fontes = []
    bruto = {}
    with tempfile.TemporaryDirectory(prefix="mini_bi_") as tmp:
        for tipo, caminho in arquivos.items():
            mod = dt.datetime.fromtimestamp(caminho.stat().st_mtime)
            log.info("Arquivo %-16s %s (modificado em %s)", tipo + ":", caminho.name, mod.strftime("%d/%m/%Y %H:%M"))
            copia = copiar_para_temporario(caminho, Path(tmp))
            cab, linhas, abas = ler_aba(copia, cfg["abas"][tipo], cfg.get("linhas_vazias_para_parar_leitura", 200))
            log.info("  aba '%s': %d linha(s) com conteúdo, %d coluna(s) no cabeçalho", cfg["abas"][tipo], len(linhas), len(cab))
            bruto[tipo] = {
                "arquivo": caminho.name, "abas": abas,
                "cabecalho": [serializar_celula(c) for c in cab],
                "linhas": [[serializar_celula(c) for c in l] for l in linhas],
            }
            # Nas planilhas brutas as datas já viram texto ISO; transformar() aceita os dois formatos
            fontes.append({"tipo": tipo, "nome": caminho.name,
                           "caminho": str(caminho.relative_to(raiz)),
                           "modificado_em": iso(mod)})

    if args.dump_bruto:
        Path(args.dump_bruto).write_text(json.dumps(bruto, ensure_ascii=False), encoding="utf-8")
        log.info("Dados brutos gravados em %s", args.dump_bruto)

    dados = transformar(bruto, cfg)
    dados["meta"].update({
        "gerado_em": iso(dt.datetime.now()),
        "arquivos_origem": fontes,
        "versao_painel": cfg.get("versao_painel"),
        "modo": "script",
    })
    dados["meta"]["validacao"] = validar(dados, cfg)

    m, r = dados["meta"], dados["meta"]["resumo"]
    log.info("Linhas lidas: %d · válidas: %d · leads no painel: %d", m["total_linhas_lidas"], m["total_linhas_validas"], len(dados["leads"]))
    log.info("Junção: %d por empresa+data, %d só por empresa, %d só no status, %d só no potencial",
             r["casados_empresa_data"], r["casados_so_empresa"], r["so_status"], r["so_potencial"])
    log.info("Meses de pedidos: %s · mês parcial: %s", ", ".join(m["meses_pedidos"]), m["mes_parcial"])
    for a in m["avisos"]:
        log.warning("AVISO: %s", a)
    nm = dados["nao_mapeados"]
    for s in nm["status"]:
        log.warning("Status não mapeado (%dx): %s", s["contagem"], s["valor"])
    for it in nm["itens"]:
        log.warning("Item para revisar (%dx): %s → %s", it["contagem"], it["item"], it["sugestao"])
    for u in nm["ufs"]:
        log.warning("UF não reconhecida (%dx): %s", u["contagem"], u["valor"])
    tipos = Counter(q["tipo_problema"] for q in dados["qualidade"])
    for t, n in tipos.most_common():
        log.info("Qualidade: %-58s %d", t, n)
    divergencias = [v for v in m["validacao"] if not v["ok"]]
    log.info("Validação contra valores de referência: %d OK, %d diferente(s)",
             len(m["validacao"]) - len(divergencias), len(divergencias))
    for v in divergencias:
        log.info("  DIFERENTE %-34s esperado %s · obtido %s (normal se a planilha mudou desde 09/10/2026)",
                 v["indicador"], v["esperado"], v["obtido"])

    html = montar_html(dados, cfg)
    pasta_saida.mkdir(parents=True, exist_ok=True)
    gravar_atomico(pasta_saida / "dados.json", json.dumps(dados, ensure_ascii=False, indent=1))
    gravar_atomico(pasta_saida / "dashboard.html", html)
    log.info("Painel gerado: %s (%.1f MB) em %.1f s", pasta_saida / "dashboard.html", len(html.encode("utf-8")) / 1e6, time.time() - inicio)
    log.info("Log: %s", destino_log)
    return 0


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Atualiza o Mini BI de Pré-Vendas da Chamfer.")
    ap.add_argument("--config", default=str(PASTA_APP / "config.json"))
    ap.add_argument("--pasta-raiz", help="Sobrescreve pasta_raiz do config (testes).")
    ap.add_argument("--saida", help="Sobrescreve pasta_de_saida do config (testes).")
    ap.add_argument("--dump-bruto", help="Grava as linhas brutas lidas em JSON (diagnóstico).")
    args = ap.parse_args()
    try:
        return executar(args)
    except ErroAtualizacao as e:
        if not log.handlers:
            configurar_log(PASTA_APP)
        log.error("ERRO: %s", e)
        log.error("O dashboard anterior foi mantido sem alterações.")
        return 1
    except ImportError as e:
        print(f"ERRO: biblioteca Python ausente ({e.name}). Rode: python -m pip install pandas openpyxl")
        return 2
    except Exception:  # noqa: BLE001 — registrar qualquer falha inesperada sem apagar o painel
        if not log.handlers:
            configurar_log(PASTA_APP)
        log.exception("ERRO inesperado. O dashboard anterior foi mantido sem alterações.")
        return 3


if __name__ == "__main__":
    sys.exit(main())
